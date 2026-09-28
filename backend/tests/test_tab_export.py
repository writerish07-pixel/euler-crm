"""Tab-scoped Excel export: current register + column/filter picker."""
import io
import os
import sys

import pytest
import pytest_asyncio

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
os.environ.setdefault("MONGO_URL", "mongodb://localhost:27017")
os.environ.setdefault("DB_NAME", "tab_export_tests")
os.environ.setdefault("JWT_SECRET", "tab-export-secret-32chars!!")

import motor.motor_asyncio  # noqa: E402
from mongomock_motor import AsyncMongoMockClient  # noqa: E402

motor.motor_asyncio.AsyncIOMotorClient = AsyncMongoMockClient

import httpx  # noqa: E402
import openpyxl  # noqa: E402
import tab_export  # noqa: E402
import server  # noqa: E402


@pytest_asyncio.fixture
async def client():
    await server.startup()
    transport = httpx.ASGITransport(app=server.app)
    async with httpx.AsyncClient(transport=transport, base_url="http://test") as c:
        r = await c.post("/api/auth/login",
                         json={"email": "owner@euler.com", "password": "euler@123"})
        c.headers.update({"Authorization": f"Bearer {r.json()['token']}"})
        yield c


def _book(content: bytes):
    return openpyxl.load_workbook(io.BytesIO(content))


@pytest.mark.asyncio
async def test_catalog_lists_registers_and_maps_paths(client):
    r = await client.get("/api/export/catalog")
    assert r.status_code == 200, r.text
    body = r.json()
    keys = {t["key"] for t in body["tabs"]}
    assert {"leads", "bookings", "payments", "finance", "claims"} <= keys
    assert body["pathToTab"]["/bookings"] == "bookings"
    assert body["pathToTab"]["/leads"] == "leads"
    leads = next(t for t in body["tabs"] if t["key"] == "leads")
    assert leads["hasPeriod"] is True
    assert "Booked" in leads["statuses"]
    assert any(c["key"] == "customerName" and c["default"] for c in leads["columns"])


@pytest.mark.asyncio
async def test_tab_export_is_only_that_register(client):
    await server.db.leads.insert_one({
        "leadId": "LD26EXP01", "customerName": "Export Cust",
        "mobile": "9000000001", "interestedModel": "Turbo Max",
        "variant": "Maxx (PV)", "executive": "Amit", "currentStatus": "Booked",
        "customerPayable": 185000, "totalReceived": 5000,
        "customerOutstanding": 180000, "bookingDate": "2026-09-10",
        "createdDate": "2026-09-01",
    })
    await server.db.bookings.insert_one({
        "bookingId": "BK26EXP01", "leadId": "LD26EXP01",
        "customerName": "Export Cust", "model": "Turbo Max",
        "variant": "Maxx (PV)", "bookingDate": "2026-09-10",
        "bookingAmount": 5000, "paymentMode": "UPI",
        "paymentReference": "UTR1", "bookingStatus": "Active",
    })
    all_wb = _book((await client.get("/api/export")).content)
    assert "Leads" in all_wb.sheetnames and "Bookings" in all_wb.sheetnames

    book = await client.get("/api/export", params={
        "tab": "bookings", "columns": "bookingId,customerName,bookingAmount",
    })
    assert book.status_code == 200, book.text
    wb = _book(book.content)
    assert wb.sheetnames == ["Booking Register"]
    rows = list(wb.active.iter_rows(values_only=True))
    assert rows[0] == ("Booking ID", "Customer", "Advance")
    assert any(r[0] == "BK26EXP01" for r in rows[1:])

    booked = await client.get("/api/export", params={
        "tab": "leads", "status": "Booked", "columns": "leadId,customerName,currentStatus",
    })
    assert booked.status_code == 200
    lead_rows = list(_book(booked.content).active.iter_rows(values_only=True))
    assert lead_rows[0] == ("Lead ID", "Customer", "Status")
    assert any(r[0] == "LD26EXP01" and r[2] == "Booked" for r in lead_rows[1:])

    empty = await client.get("/api/export", params={
        "tab": "leads", "status": "Lost", "columns": "leadId",
    })
    lost_rows = list(_book(empty.content).active.iter_rows(values_only=True))
    assert lost_rows[0] == ("Lead ID",)
    assert not any(r[0] == "LD26EXP01" for r in lost_rows[1:])


@pytest.mark.asyncio
async def test_tab_export_search_and_period(client):
    await server.db.payments.insert_many([
        {"receiptNumber": "RC-A", "leadId": "LD1", "customerName": "Ramesh",
         "date": "2026-09-12", "amount": 1000, "paymentMode": "Cash"},
        {"receiptNumber": "RC-B", "leadId": "LD2", "customerName": "Suresh",
         "date": "2026-08-12", "amount": 2000, "paymentMode": "UPI"},
    ])
    r = await client.get("/api/export", params={
        "tab": "payments", "month": "2026-09", "q": "ramesh",
        "columns": "receiptNumber,customerName,date",
    })
    assert r.status_code == 200, r.text
    rows = list(_book(r.content).active.iter_rows(values_only=True))
    data = [row[0] for row in rows[1:]]
    assert "RC-A" in data
    assert "RC-B" not in data


@pytest.mark.asyncio
async def test_unknown_tab_is_422(client):
    r = await client.get("/api/export", params={"tab": "not-a-register"})
    assert r.status_code == 422


@pytest.mark.asyncio
async def test_executive_cannot_export(client):
    transport = httpx.ASGITransport(app=server.app)
    async with httpx.AsyncClient(transport=transport, base_url="http://test") as c:
        r = await c.post("/api/auth/login",
                         json={"email": "executive@euler.com", "password": "euler@123"})
        c.headers.update({"Authorization": f"Bearer {r.json()['token']}"})
        assert (await c.get("/api/export", params={"tab": "leads"})).status_code == 403
        assert (await c.get("/api/export/catalog")).status_code == 403


def test_tab_for_path():
    assert tab_export.tab_for_path("/bookings") == "bookings"
    assert tab_export.tab_for_path("/leads") == "leads"
    assert tab_export.tab_for_path("/oem-claims/no-vehicle") == "oem_claims"
    assert tab_export.tab_for_path("/") == "leads"
    assert tab_export.tab_for_path("/settings") == "leads"
